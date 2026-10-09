"""
Recipe R6 — weekly Excel report about a uCoz blog, built in Python with openpyxl and run from cron.

Why Python here: openpyxl writes real .xlsx files (styles, number formats, filters, charts, links) — the kind
of document PHP has no built-in library for. The job runs from cron as a plain script, so it is not bound by
the ~60 s HTTP limit.

App root: domains/{domain}/app_python/{name}/  (next to app.py and the passenger_wsgi.py stub)
Packages: requirements.txt (openpyxl, requests) installed by the one-off cron pip job (PYTHON.md §5).
Config:   config.json, mode 0600:
            {"site": "blog.example.com", "uapi_token": "<read-only uAPI key>", "staff_gids": [4],
             "tg_token": "<bot token, optional>", "tg_chat": 123456789}
Cron:     minute=0 hour=6 dayofweek=1  (Monday 06:00, server time)
          command="cd /home/{user}/domains/{domain}/app_python/{name} &&
                   /home/{user}/virtualenv/domains/{domain}/app_python/{name}/3.11/bin/python report.py >> logs/cron.log 2>&1"
Output:   data/reports/blog-report-YYYY-MM-DD.xlsx (last 12 kept), data/last.json; optional Telegram sendDocument.
Usage:    python report.py            build + send
          python report.py --no-send  build only
"""
import datetime as dt
import html
import json
import os
import sys
import time
import urllib.parse
import urllib.request

DIR = os.path.dirname(os.path.abspath(__file__))
REPORTS = os.path.join(DIR, "data", "reports")
STATE = os.path.join(DIR, "data", "last.json")
LOG = os.path.join(DIR, "logs", "report.log")
KEEP = 12
BLUE = "1D6FE8"


def log(event, **ctx):
    os.makedirs(os.path.dirname(LOG), exist_ok=True)
    with open(LOG, "a", encoding="utf-8") as f:
        f.write(f"{time.strftime('%Y-%m-%dT%H:%M:%S')} {event} {json.dumps(ctx, ensure_ascii=False)}\n")


def load_config():
    with open(os.path.join(DIR, "config.json"), encoding="utf-8") as f:
        return json.load(f)


# ---------- uAPI ----------
def uapi(cfg, path, query=None):
    url = f"https://{cfg['site']}/uapi/{path.lstrip('/')}"
    if query:
        url += "?" + urllib.parse.urlencode(query)
    req = urllib.request.Request(url, headers={"Authorization": "Bearer " + cfg["uapi_token"], "Accept": "application/json"})
    with urllib.request.urlopen(req, timeout=30) as r:
        data = json.loads(r.read().decode("utf-8"))
    if isinstance(data, dict) and data.get("error"):
        err = data["error"]
        raise RuntimeError("uapi_" + str(err.get("code") if isinstance(err, dict) else err))
    return data


def text(s):
    return html.unescape(str(s or "")).strip()


def fetch_posts(cfg):
    posts, page, pages = [], 1, 1
    while page <= pages:
        r = uapi(cfg, "blog", {"per_page": 100, "page": page})   # 100 = maximum for /blog
        pages = int(r.get("pages") or 1)
        for p in r.get("blogs") or []:
            if isinstance(p, dict) and p.get("id"):
                posts.append(p)
        page += 1
    return posts


def fetch_categories(cfg):
    r = uapi(cfg, "blog/category")
    items = r.values() if isinstance(r, dict) else r
    return {int(c["id"]): text(c.get("name")) for c in items if isinstance(c, dict) and c.get("id")}


def normalize(posts, cats, today):
    rows = []
    for p in posts:
        cat = p.get("category") or {}
        cid = int(cat.get("id") or 0) if isinstance(cat, dict) else 0
        added = dt.datetime.fromtimestamp(int(p.get("add_date_ts") or 0), dt.timezone.utc).date()
        views = max(0, int(p.get("reads") or 0))
        rows.append({
            "id": int(p["id"]),
            "title": text(p.get("title")),
            "url": str(p.get("entry_url") or ""),
            "category": cats.get(cid) or (text(cat.get("name")) if isinstance(cat, dict) else "") or "—",
            "date": added,
            "views": views,
            "comments": max(0, int(p.get("comments_num") or 0)),
            "per_day": round(views / max(1, (today - added).days), 1),
        })
    rows.sort(key=lambda r: r["views"], reverse=True)   # never trust the list order
    return rows


# ---------- workbook ----------
def build_xlsx(rows, site, path, today):
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.formatting.rule import DataBarRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    head_font, head_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor=BLUE)
    link_font = Font(color=BLUE, underline="single")

    def table(ws, headers, widths, data, formats=None):
        ws.append(headers)
        for cell in ws[1]:
            cell.font, cell.fill = head_font, head_fill
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[1].height = 30
        for r in data:
            ws.append(r)
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for col, fmt in (formats or {}).items():
            for (cell,) in ws.iter_rows(min_row=2, min_col=col, max_col=col):
                cell.number_format = fmt
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    wb = Workbook()

    # Summary
    ws = wb.active
    ws.title = "Summary"
    total_views = sum(r["views"] for r in rows)
    recent = [r for r in rows if (today - r["date"]).days <= 30]
    top = rows[0] if rows else None
    ws["A1"] = f"Blog report for {site}"
    ws["A1"].font = Font(bold=True, size=16)
    ws["A2"] = f"Generated {today.isoformat()}"
    ws["A2"].font = Font(color="808080")
    kpis = [
        ("Posts", len(rows)),
        ("Views, all time", total_views),
        ("Comments", sum(r["comments"] for r in rows)),
        ("Average views per post", round(total_views / max(1, len(rows)))),
        ("New posts in 30 days", len(recent)),
    ]
    for i, (label, value) in enumerate(kpis, start=4):
        ws.cell(row=i, column=1, value=label)
        c = ws.cell(row=i, column=2, value=value)
        c.number_format, c.font = "#,##0", Font(bold=True, size=12)
    if top:
        ws.cell(row=10, column=1, value="Most read post")
        c = ws.cell(row=10, column=2, value=top["title"])
        c.hyperlink, c.font = top["url"] or None, link_font
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 34, 70

    # Posts
    ws = wb.create_sheet("Posts")
    table(ws, ["#", "Post", "Category", "Date", "Views", "Comments", "Views per day"],
          [6, 70, 22, 12, 13, 13, 12],
          [[i, r["title"], r["category"], r["date"], r["views"], r["comments"], r["per_day"]] for i, r in enumerate(rows, 1)],
          {4: "YYYY-MM-DD", 5: "#,##0", 6: "#,##0", 7: "#,##0.0"})
    for i, r in enumerate(rows, 2):
        if r["url"]:
            ws.cell(row=i, column=2).hyperlink = r["url"]
            ws.cell(row=i, column=2).font = link_font
    if rows:
        ws.conditional_formatting.add(f"E2:E{len(rows) + 1}", DataBarRule(start_type="min", end_type="max", color=BLUE))

    # Categories
    by_cat = {}
    for r in rows:
        c = by_cat.setdefault(r["category"], {"n": 0, "v": 0, "c": 0, "last": r["date"]})
        c["n"] += 1
        c["v"] += r["views"]
        c["c"] += r["comments"]
        c["last"] = max(c["last"], r["date"])
    cats = sorted(by_cat.items(), key=lambda kv: kv[1]["v"], reverse=True)
    ws = wb.create_sheet("Categories")
    table(ws, ["Category", "Posts", "Views", "Average per post", "Comments", "Latest post"],
          [26, 10, 14, 14, 14, 14],
          [[k, v["n"], v["v"], round(v["v"] / v["n"]), v["c"], v["last"]] for k, v in cats],
          {3: "#,##0", 4: "#,##0", 5: "#,##0", 6: "YYYY-MM-DD"})
    if cats:
        ch = BarChart()
        ch.type, ch.style, ch.title, ch.legend = "bar", 10, "Average views per post", None
        ch.add_data(Reference(ws, min_col=4, min_row=1, max_row=len(cats) + 1), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=len(cats) + 1))
        ch.height, ch.width = max(7.5, 0.6 * len(cats)), 18
        ws.add_chart(ch, "H2")

    # Years
    by_year = {}
    for r in rows:
        y = by_year.setdefault(r["date"].year, {"n": 0, "v": 0})
        y["n"] += 1
        y["v"] += r["views"]
    years = sorted(by_year.items())
    ws = wb.create_sheet("By year")
    table(ws, ["Year", "Posts", "Views"], [8, 10, 14], [[y, v["n"], v["v"]] for y, v in years], {3: "#,##0"})
    if years:
        ch = BarChart()
        ch.style, ch.title, ch.legend = 10, "Post views by publication year", None
        ch.add_data(Reference(ws, min_col=3, min_row=1, max_row=len(years) + 1), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=len(years) + 1))
        ch.height, ch.width = 8, 18
        ws.add_chart(ch, "E2")

    tmp = path + ".tmp"
    wb.save(tmp)
    os.replace(tmp, path)   # readers never see a half-written file


# ---------- delivery ----------
def send_telegram(cfg, path, caption):
    if not cfg.get("tg_token") or not cfg.get("tg_chat"):
        return False
    import requests   # installed by the cron pip job
    with open(path, "rb") as f:
        r = requests.post(f"https://api.telegram.org/bot{cfg['tg_token']}/sendDocument",
                          data={"chat_id": cfg["tg_chat"], "caption": caption},
                          files={"document": (os.path.basename(path), f)}, timeout=30)
    ok = r.ok and r.json().get("ok")
    if not ok:
        raise RuntimeError("tg: " + str(r.json().get("description", r.status_code))[:120])   # never log the URL
    return True


def prune():
    files = sorted(f for f in os.listdir(REPORTS) if f.endswith(".xlsx"))
    for f in files[:-KEEP]:
        os.remove(os.path.join(REPORTS, f))


def main(send=True):
    t0 = time.time()
    cfg = load_config()
    today = dt.date.today()
    os.makedirs(REPORTS, exist_ok=True)
    rows = normalize(fetch_posts(cfg), fetch_categories(cfg), today)
    if not rows:
        raise RuntimeError("no_posts")
    name = f"blog-report-{today.isoformat()}.xlsx"
    path = os.path.join(REPORTS, name)
    build_xlsx(rows, cfg["site"], path, today)
    views = f"{sum(r['views'] for r in rows):,}"   # 14,600,000
    caption = f"Blog report for {cfg['site']}, {today.isoformat()}: {len(rows)} posts, {views} views"
    sent = send_telegram(cfg, path, caption) if send else False
    prune()
    state = {"at": time.strftime("%Y-%m-%dT%H:%M:%S"), "file": name, "posts": len(rows),
             "bytes": os.path.getsize(path), "seconds": round(time.time() - t0, 1), "sent": sent}
    with open(STATE + ".tmp", "w", encoding="utf-8") as f:
        json.dump(state, f)
    os.replace(STATE + ".tmp", STATE)
    log("report", **state)
    return state


if __name__ == "__main__":
    try:
        print(json.dumps(main(send="--no-send" not in sys.argv), ensure_ascii=False))
    except Exception as e:   # cron output goes to logs/cron.log; keep a line in our own log too
        log("report_failed", error=str(e)[:300])
        raise
